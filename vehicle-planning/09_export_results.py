#!/usr/bin/env python
import argparse
import datetime
import importlib.util
import os
import sys
import zipfile

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))


def mod(fname, name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fname))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


summ = mod("03_summarize.py", "summ")
merge = mod("07_merge_chunks.py", "merge")
speed = mod("08_speed_analysis.py", "speed")

PAIRS = [("fd", "yflow-kin-v2"), ("fd", "yflow-corr-v2"), ("fd", "yflow-kin"), ("fd", "yflow-corr"),
         ("fdstar", "yflow_star-kin-v2"), ("fdstar", "yflow_star-corr-v2"), ("fd", "fdstar"), ("fd", "yflow_star-kin-v2"),
         ("fd", "fd_clip"), ("fd", "fd_post"), ("fd_clip", "yflow-kin-v2"), ("fd_post", "yflow-kin-v2"),
         ("fd_post", "yflow_post-kin-v2"), ("yflow-kin-v2", "yflow_post-kin-v2")]


def all_runs(exp_root, results_dir):
    import glob
    rows = []
    for p in sorted(glob.glob(os.path.join(exp_root, "**", "aggregator_metric", "*.parquet"), recursive=True)):
        df = pd.read_parquet(p)
        fin = df[df["scenario"] == "final_score"] if "scenario" in df else df.tail(1)
        if fin.empty:
            continue
        fin = fin.iloc[0]
        uid = summ.uid_from_path(p, exp_root)
        parts = uid.split("/")
        row = {"run": uid, "split": parts[1] if len(parts) > 1 else "", "ch": parts[2] if len(parts) > 2 else "",
               "tag": parts[3] if len(parts) > 3 else "", "stamp": parts[4] if len(parts) > 4 else "",
               "n_scen": int(fin["num_scenarios"]) if "num_scenarios" in fin and pd.notna(fin["num_scenarios"]) else -1}
        for col, name in summ.METRICS:
            if col in fin and pd.notna(fin[col]):
                row[name] = round(100 * float(fin[col]), 2) if col == "score" else round(float(fin[col]), 4)
        row.update({k: (round(v, 3) if isinstance(v, float) else v) for k, v in summ.yflow_stats(uid, results_dir).items()})
        rows.append(row)
    return pd.DataFrame(rows)


def merged(exp_root, results_dir, data_root, split):
    import json
    runs = merge.find_runs(exp_root, split)
    st_path = os.path.join(data_root, "chunks", split, "state.json")
    expected = json.load(open(st_path))["n_tokens"] if os.path.exists(st_path) else None
    table = []
    for (group, ch), by_id in sorted(runs.items()):
        frames = [merge.scenario_rows(pd.read_parquet(p)) for p, _ in by_id.values()]
        rows = pd.concat(frames).drop_duplicates(subset="scenario", keep="last")
        row = {"method": group, "ch": ch, "chunks": ",".join(str(k) for k in sorted(by_id))}
        row.update(merge.aggregate(rows))
        if expected:
            row["coverage"] = f"{row['n_scen']}/{expected}"
        row.update(merge.yflow_summary([u for _, u in by_id.values()], results_dir))
        table.append(row)
    return pd.DataFrame(table)


def pairs_table(mdf):
    out = []
    for ch in sorted(mdf["ch"].unique()):
        cls = mdf[mdf["ch"] == ch].set_index("method")["CLS"]
        for a, b in PAIRS:
            if a in cls and b in cls:
                out.append({"ch": ch, "비교": f"{a} → {b}", "기준 CLS": cls[a], "대상 CLS": cls[b],
                            "Δ CLS": round(cls[b] - cls[a], 2)})
    return pd.DataFrame(out)


def speed_sheets(exp_root, split, ch, methods, base):
    data = {}
    for t in methods:
        try:
            data[t], _ = speed.load(exp_root, split, ch, t)
        except SystemExit:
            pass
    if base not in data or len(data) < 2:
        return {}
    common = sorted(set.intersection(*[set(d.index) for d in data.values()]))
    data = {t: d.loc[common] for t, d in data.items()}
    b = data[base]
    viol = b.index[(b["n_viol"] > 0) | (b["speed_limit_compliance"] < 1.0)]
    rest = b.index.difference(viol)
    blocks = []
    s1 = pd.DataFrame([speed.summary_row(t, d) for t, d in data.items()]); s1.insert(0, "구분", "[1] 전체")
    blocks.append(s1)
    if len(viol):
        s2 = pd.DataFrame([speed.summary_row(t, d.loc[viol]) for t, d in data.items()])
        s2.insert(0, "구분", f"[2] {base}가 과속한 {len(viol)}개")
        blocks.append(s2)
        s2b = pd.DataFrame([{"method": t, "n": len(rest), "CLS": round(100 * d.loc[rest, "score"].mean(), 2)} for t, d in data.items()])
        s2b.insert(0, "구분", f"[2b] 과속 없던 {len(rest)}개")
        blocks.append(s2b)
    dec = pd.DataFrame({t: speed.decompose(d) for t, d in data.items()}).T.reset_index().rename(columns={"index": "method"})
    dec.insert(0, "구분", "[4] CLS 분해")
    blocks.append(dec)
    sp = pd.concat(blocks, ignore_index=True)

    rows = []
    for stype, idx in b.groupby("scenario_type").groups.items():
        r = {"scenario_type": stype, "n": len(idx), f"{base} 위반 %": round(100 * (b.loc[idx, "n_viol"] > 0).mean(), 1)}
        for t, d in data.items():
            r[f"speed {t}"] = round(d.loc[idx, "speed_limit_compliance"].mean(), 4)
        for t, d in data.items():
            r[f"CLS {t}"] = round(100 * d.loc[idx, "score"].mean(), 2)
        rows.append(r)
    types = pd.DataFrame(rows).sort_values(f"{base} 위반 %", ascending=False)

    cols = ["scenario_type", "score", "speed_limit_compliance", "n_viol", "max_over",
            "ego_progress_along_expert_route", "no_ego_at_fault_collisions", "time_to_collision_within_bound"]
    scen = pd.concat({t: d[[c for c in cols if c in d]] for t, d in data.items()}, axis=1)
    scen.columns = [f"{t} | {c}" for t, c in scen.columns]
    scen = scen.reset_index().rename(columns={"index": "scenario"})
    return {"speed": sp, "types": types, "scen": scen}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--exp_root", default=os.environ.get("NUPLAN_EXP_ROOT", "/root/nuplan/exp"))
    ap.add_argument("--results_dir", default=os.environ.get("RESULTS_DIR", "/root/fd_yflow_results"))
    ap.add_argument("--data_root", default=os.environ.get("NUPLAN_DATA_ROOT", "/root/nuplan/dataset"))
    ap.add_argument("--out", default="")
    a = ap.parse_args()
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    out = a.out or os.path.join(a.results_dir, f"results_{stamp}.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)

    sheets = {}
    ar = all_runs(a.exp_root, a.results_dir)
    sheets["all_runs"] = ar
    print(f"[export] all_runs: {len(ar)}개 실행")

    for split in ("test14-hard", "val14", "test14-random"):
        mdf = merged(a.exp_root, a.results_dir, a.data_root, split)
        if mdf.empty:
            continue
        sheets[f"{split}_merged"] = mdf
        pt = pairs_table(mdf)
        if not pt.empty:
            sheets[f"{split}_pairs"] = pt
        methods = sorted({m for m in mdf["method"] if "[" not in m}, key=lambda x: (x != "fd", x))
        for ch in sorted(mdf["ch"].unique()):
            res = speed_sheets(a.exp_root, split, ch, methods, "fd")
            for k, v in res.items():
                sheets[f"{k}_{split}_{ch}"] = v
        print(f"[export] {split}: 방법 {len(mdf)}행")

    if not ar.empty:
        dev = sorted({t for t in ar.loc[(ar["split"] == "mini") & ar["tag"].str.endswith("-dev"), "tag"]})
        if "fd-dev" in dev:
            res = speed_sheets(a.exp_root, "mini", "nr", ["fd-dev"] + [t for t in dev if t != "fd-dev"], "fd-dev")
            for k, v in res.items():
                sheets[f"{k}_mini-dev"] = v

    sheets = {k[:31]: v for k, v in sheets.items()}
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        with pd.ExcelWriter(out, engine="openpyxl") as w:
            for name, df in sheets.items():
                df.to_excel(w, sheet_name=name, index=False)
                ws = w.sheets[name]
                for cell in ws[1]:
                    cell.font = Font(name="Arial", bold=True)
                    cell.fill = PatternFill("solid", fgColor="DCE6F2")
                for col in ws.columns:
                    width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                    ws.column_dimensions[col[0].column_letter].width = min(max(8, width + 2), 60)
                ws.freeze_panes = "B2"
        print(f"[export] -> {out}  (시트 {len(sheets)}개: {', '.join(sheets)})")
    except ImportError:
        zpath = os.path.splitext(out)[0] + "_csv.zip"
        with zipfile.ZipFile(zpath, "w") as z:
            for name, df in sheets.items():
                z.writestr(f"{name}.csv", df.to_csv(index=False).encode("utf-8-sig"))
        print(f"[export] openpyxl 없음 -> CSV zip: {zpath}  (엑셀로 만들려면 pip install openpyxl 후 다시 실행)")


if __name__ == "__main__":
    main()
